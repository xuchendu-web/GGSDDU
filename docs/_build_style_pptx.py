#!/usr/bin/env python3
"""公募经营驾驶舱 · 风格拆解 PPT（Action Title）"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from lxml import etree
from copy import deepcopy

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLUE = RGBColor(0x16, 0x77, 0xFF)
BLUE_DK = RGBColor(0x0F, 0x62, 0xD8)
NAVY = RGBColor(0x1F, 0x23, 0x29)
GRAY = RGBColor(0x4E, 0x59, 0x69)
GRAY2 = RGBColor(0x86, 0x90, 0x9C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xEE, 0xF3, 0xF9)
RED = RGBColor(0xE5, 0x48, 0x4D)
GREEN = RGBColor(0x2B, 0xA4, 0x71)
PURPLE = RGBColor(0x8B, 0x6C, 0xFF)
ORANGE = RGBColor(0xFA, 0x8C, 0x16)
CARD = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, size=14, bold=False, color=NAVY, name="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", name)


def add_text(box, text, size=14, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return tf


def add_para(tf, text, size=13, bold=False, color=GRAY, space=6):
    p = tf.add_paragraph()
    p.space_before = Pt(space)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return p


def bg(slide, color=BG):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s


def bar(slide, l, t, w, h, color=BLUE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.adjustments[0] = 0.15
    return s


def card(slide, l, t, w, h):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = WHITE
    s.line.color.rgb = RGBColor(0xE6, 0xEE, 0xF6)
    s.adjustments[0] = 0.06
    return s


def footer(slide, n, total=8):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(10), Inches(0.25))
    add_text(box, "经营驾驶舱风格拆解  ·  演示口径，产品已匿名", 10, False, GRAY2)
    box2 = slide.shapes.add_textbox(Inches(11.6), Inches(7.15), Inches(1.2), Inches(0.25))
    add_text(box2, f"{n} / {total}", 10, False, GRAY2, PP_ALIGN.RIGHT)


# ===== 1 cover =====
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s, RGBColor(0x0F, 0x3D, 0x7A))
bar(s, Inches(0.7), Inches(2.15), Inches(0.12), Inches(1.6), BLUE)
tb = s.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(11), Inches(0.9))
add_text(tb, "白底 Bento + 金融蓝锚点：这是一套给管理层扫读、给投研下钻的公募经营驾驶舱", 26, True, WHITE)
tb = s.shapes.add_textbox(Inches(1.05), Inches(3.15), Inches(10), Inches(0.8))
add_text(tb, "风格拆解  ·  交互蓝图  ·  可运行 HTML 原型", 16, False, RGBColor(0xBA, 0xE0, 0xFF))
tb = s.shapes.add_textbox(Inches(1.05), Inches(6.4), Inches(10), Inches(0.35))
add_text(tb, "内部工作台语言  |  红涨绿跌  |  卡片内过滤  |  数据截至 T+1", 12, False, RGBColor(0x9A, 0xC4, 0xF0))

# ===== 2 定位 =====
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
bar(s, Inches(0.5), Inches(0.38), Inches(0.1), Inches(0.42))
tb = s.shapes.add_textbox(Inches(0.72), Inches(0.32), Inches(12), Inches(0.55))
add_text(tb, "它不是理财 App，也不是彭博黑底终端——而是机构内部的「一屏看全公司」", 20, True, BLUE_DK)

items = [
    ("像谁", "Ant Design Pro / 券商内部投研台：浅冷色画布、白卡片、蓝竖条标题、等宽 KPI。"),
    ("不像谁", "消费级渐变插画、暗色指挥大屏、3D 饼图、游戏化血条。风控半环只表达时长结构。"),
    ("核心承诺", "30 秒看规模/申赎/营收是否异常；再一层下到产品表与研究员组合。"),
    ("密度策略", "高密度但可扫读：主数字 26–28px，标签 11–12px 灰色，图表去描边、图例置顶。"),
]
for i, (k, v) in enumerate(items):
    y = 1.15 + i * 1.35
    card(s, Inches(0.5), Inches(y), Inches(12.3), Inches(1.2))
    bar(s, Inches(0.7), Inches(y + 0.38), Inches(0.08), Inches(0.42), BLUE if i != 1 else ORANGE)
    tb = s.shapes.add_textbox(Inches(1.0), Inches(y + 0.18), Inches(11.4), Inches(0.35))
    add_text(tb, k, 14, True, BLUE_DK if i != 1 else ORANGE)
    tb = s.shapes.add_textbox(Inches(1.0), Inches(y + 0.55), Inches(11.4), Inches(0.5))
    add_text(tb, v, 14, False, GRAY)
footer(s, 2)

# ===== 3 IA =====
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
bar(s, Inches(0.5), Inches(0.38), Inches(0.1), Inches(0.42))
tb = s.shapes.add_textbox(Inches(0.72), Inches(0.32), Inches(12), Inches(0.55))
add_text(tb, "四层信息架构：经营总览 → 公司全景 → 投资切片 → 研究绩效，禁止混层", 20, True, BLUE_DK)

layers = [
    ("L0 经营管理", "只数 / 规模 / 生命周期 / 当日净申购 / 管理费与盈亏", "此刻该不该紧张"),
    ("L1 公司全景", "同业雷达 · 规模结构 · 产品榜 · 管理费节奏 · 渠道散点", "我们在行业里站哪"),
    ("L2 基金投资", "部门/经理下拉 + 产品明细表 + 静态风控", "谁的组合、什么风险"),
    ("L3 权益研究", "考评榜 · 组合 vs 基准 · 行业收益条", "观点有没有被市场验证"),
]
for i, (t, d, q) in enumerate(layers):
    x = 0.5 + i * 3.2
    card(s, Inches(x), Inches(1.3), Inches(3.0), Inches(5.2))
    bar(s, Inches(x), Inches(1.3), Inches(3.0), Inches(0.12))
    tb = s.shapes.add_textbox(Inches(x + 0.18), Inches(1.6), Inches(2.65), Inches(0.8))
    add_text(tb, t, 16, True, BLUE_DK)
    tb = s.shapes.add_textbox(Inches(x + 0.18), Inches(2.5), Inches(2.65), Inches(2.2))
    add_text(tb, d, 13, False, GRAY)
    tb = s.shapes.add_textbox(Inches(x + 0.18), Inches(5.2), Inches(2.65), Inches(0.9))
    add_text(tb, q, 12, True, BLUE)
footer(s, 3)

# ===== 4 视觉系统 =====
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
bar(s, Inches(0.5), Inches(0.38), Inches(0.1), Inches(0.42))
tb = s.shapes.add_textbox(Inches(0.72), Inches(0.32), Inches(12), Inches(0.55))
add_text(tb, "视觉系统可复用：画布冷灰、卡片纯白、语义色固定，规模堆叠禁止「彩虹重排」", 20, True, BLUE_DK)

swatches = [
    ("#EEF3F9", "画布", RGBColor(0xEE, 0xF3, 0xF9), NAVY),
    ("#1677FF", "品牌蓝", BLUE, WHITE),
    ("#1F2329", "主数字", NAVY, WHITE),
    ("#E5484D", "涨/警示", RED, WHITE),
    ("#2BA471", "跌/对照", GREEN, WHITE),
    ("#8B6CFF", "生命周期", PURPLE, WHITE),
    ("#FA8C16", "规模/营收", ORANGE, WHITE),
    ("#5B7CFA", "固收堆叠", RGBColor(0x5B, 0x7C, 0xFA), WHITE),
]
for i, (hexv, name, c, tc) in enumerate(swatches):
    x = 0.5 + (i % 8) * 1.58
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.2), Inches(1.45), Inches(1.55))
    sh.fill.solid(); sh.fill.fore_color.rgb = c; sh.line.fill.background(); sh.adjustments[0] = 0.12
    tb = s.shapes.add_textbox(Inches(x + 0.06), Inches(1.35), Inches(1.32), Inches(0.7))
    add_text(tb, name, 12, True, tc, PP_ALIGN.CENTER)
    tb = s.shapes.add_textbox(Inches(x + 0.06), Inches(2.15), Inches(1.32), Inches(0.4))
    add_text(tb, hexv, 10, False, tc, PP_ALIGN.CENTER)

rules = [
    "标题：3×14 蓝竖条 + 14px/650，不放花哨图标字体。",
    "KPI：26–28px 等宽；副文案 11px 灰。数量用蓝玻璃图标，金额用橙。",
    "A 股口径红涨绿跌；风控最长天数用红，短违规用蓝。",
    "卡片：圆角 10、间距 10、边框 #E9F0F6、投影极轻。默认不做暗色。",
]
for i, r in enumerate(rules):
    y = 3.1 + i * 0.85
    card(s, Inches(0.5), Inches(y), Inches(12.3), Inches(0.75))
    tb = s.shapes.add_textbox(Inches(0.75), Inches(y + 0.18), Inches(11.8), Inches(0.45))
    add_text(tb, r, 15, False, NAVY)
footer(s, 4)

# ===== 5 图表 =====
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
bar(s, Inches(0.5), Inches(0.38), Inches(0.1), Inches(0.42))
tb = s.shapes.add_textbox(Inches(0.72), Inches(0.32), Inches(12), Inches(0.55))
add_text(tb, "图表有固定语义：气泡看生命周期，堆叠看结构，对数散点看渠道质量", 20, True, BLUE_DK)

charts = [
    ("气泡图", "产品形态", "X=阶段，Y 与半径=只数。发行期为 0 留空，不画假点。"),
    ("胶囊柱", "净申购", "零轴上下对称，红入绿出。当日 KPI 用偏赭红，避免与涨跌红撞车。"),
    ("雷达", "同业结构", "六品类；排名/市占放左侧大字，不塞进雷达中心。"),
    ("堆叠面积", "规模走势", "无描边；固收/货币做主体，饱和度压住。快捷档改写区间。"),
    ("柱+线", "管理费", "去年蓝柱、今年橙柱，累计用线。未发生月份 null，不补 0。"),
    ("对数散点", "销售渠道", "规模跨数量级必须 log。银行深蓝 / 券商蓝 / 三方绿。"),
]
for i, (k, w, d) in enumerate(charts):
    r, c = divmod(i, 3)
    x, y = 0.5 + c * 4.2, 1.2 + r * 2.7
    card(s, Inches(x), Inches(y), Inches(4.0), Inches(2.5))
    tb = s.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.2), Inches(3.55), Inches(0.4))
    add_text(tb, k, 16, True, BLUE_DK)
    tb = s.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.65), Inches(3.55), Inches(0.35))
    add_text(tb, w, 12, True, BLUE)
    tb = s.shapes.add_textbox(Inches(x + 0.22), Inches(y + 1.1), Inches(3.55), Inches(1.1))
    add_text(tb, d, 13, False, GRAY)
footer(s, 5)

# ===== 6 交互 =====
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
bar(s, Inches(0.5), Inches(0.38), Inches(0.1), Inches(0.42))
tb = s.shapes.add_textbox(Inches(0.72), Inches(0.32), Inches(12), Inches(0.55))
add_text(tb, "交互下沉到卡片：页头只留截至日期，筛选项跟着图表走", 20, True, BLUE_DK)

left = [
    "页级：数据截至日期（T+1 时间戳）",
    "卡片级：Segmented、日期对、部门/经理下拉",
    "行级：榜单与产品表 hover/选中，驱动邻图",
    "空态：缺数用「—」，同比无数据不画箭头",
]
right = [
    "不要顶栏堆 8 个全局筛选",
    "不要把投研组合和管理费画同一张图",
    "不要在客户材料写真实产品名",
    "移动端另做 KPI 简报，不要硬缩这张屏",
]
card(s, Inches(0.5), Inches(1.2), Inches(6.0), Inches(5.3))
tb = s.shapes.add_textbox(Inches(0.75), Inches(1.4), Inches(5.5), Inches(0.45))
add_text(tb, "做", 16, True, GREEN)
tb = s.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(5.5), Inches(4.1))
tf = add_text(tb, left[0], 15, False, NAVY)
for x in left[1:]:
    add_para(tf, x, 15, False, NAVY, 14)

card(s, Inches(6.8), Inches(1.2), Inches(6.0), Inches(5.3))
tb = s.shapes.add_textbox(Inches(7.05), Inches(1.4), Inches(5.5), Inches(0.45))
add_text(tb, "不做", 16, True, RED)
tb = s.shapes.add_textbox(Inches(7.05), Inches(2.0), Inches(5.5), Inches(4.1))
tf = add_text(tb, right[0], 15, False, NAVY)
for x in right[1:]:
    add_para(tf, x, 15, False, NAVY, 14)
footer(s, 6)

# ===== 7 软件 =====
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
bar(s, Inches(0.5), Inches(0.38), Inches(0.1), Inches(0.42))
tb = s.shapes.add_textbox(Inches(0.72), Inches(0.32), Inches(12), Inches(0.55))
add_text(tb, "落地路径：先冻结 Widget 协议与口径字典，再接 T+1 估值/TA/同业披露", 20, True, BLUE_DK)

stack = [
    ("前端", "React 或 Vue 3 + CSS Grid + ECharts 5。每个卡片 = title / filters / renderer / queryKey。"),
    ("数据", "经营层：估值规模 + TA 申赎 + 管理费计提。同业：协会/年报披露。投研：模拟盘 vs 基准。"),
    ("权限", "L0–L1 管理层；L2 按投资部门；L3 研究团队。同一视觉，不同 query。"),
    ("分辨率", "先做 1920 工作台，1440 只压缩字号。移动端出简报，不做响应式大屏。"),
    ("口径纪律", "归因必须与净值勾稽；说不清标【待核实】。客户材料不出现具体产品名。"),
    ("交付物", "本仓库已含可点击 HTML 原型、设计 token、本页 PPT。下一步接真实接口。"),
]
for i, (k, v) in enumerate(stack):
    r, c = divmod(i, 2)
    x, y = 0.5 + c * 6.4, 1.2 + r * 1.75
    card(s, Inches(x), Inches(y), Inches(6.15), Inches(1.6))
    tb = s.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.18), Inches(5.7), Inches(0.4))
    add_text(tb, k, 15, True, BLUE_DK)
    tb = s.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.62), Inches(5.7), Inches(0.8))
    add_text(tb, v, 13, False, GRAY)
footer(s, 7)

# ===== 8 next =====
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
bar(s, Inches(0.5), Inches(0.38), Inches(0.1), Inches(0.42))
tb = s.shapes.add_textbox(Inches(0.72), Inches(0.32), Inches(12), Inches(0.55))
add_text(tb, "建议立刻开工的三件事：口径表、Widget 清单、权限矩阵——视觉已经够抄", 20, True, BLUE_DK)

steps = [
    ("01", "口径字典", "规模用期末还是日均、管理费含不含货基、同业是否含货币，先写成一页纸。"),
    ("02", "Widget 冻结", "按本原型 14 张卡片立项，每张写数据源、刷新频率、负责人。"),
    ("03", "权限矩阵", "管理层/固收/权益/研究看到的字段不同，表头先裁再开发。"),
]
for i, (n, t, d) in enumerate(steps):
    y = 1.25 + i * 1.7
    card(s, Inches(0.5), Inches(y), Inches(12.3), Inches(1.5))
    tb = s.shapes.add_textbox(Inches(0.75), Inches(y + 0.4), Inches(1.2), Inches(0.6))
    add_text(tb, n, 28, True, BLUE)
    tb = s.shapes.add_textbox(Inches(2.1), Inches(y + 0.25), Inches(10.3), Inches(0.4))
    add_text(tb, t, 18, True, NAVY)
    tb = s.shapes.add_textbox(Inches(2.1), Inches(y + 0.75), Inches(10.3), Inches(0.5))
    add_text(tb, d, 14, False, GRAY)
footer(s, 8)

out = "/workspace/docs/公募经营驾驶舱-风格拆解.pptx"
prs.save(out)
print("saved", out)
