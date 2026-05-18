"""Reverse-read the generated PPTX and verify slides and titles."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

PROJECT_ROOT = Path(__file__).resolve().parent.parent
PPT = PROJECT_ROOT / "output" / "FOF资产配置新思路.pptx"

EXPECTED_PAGES = 42

EXPECTED_KEYWORDS = {
    1:  ["公募FOF投资新思路"],
    2:  ["目录"],
    5:  ["大类资产长期画像"],
    10: ["监管边界"],
    12: ["相关性矩阵"],
    20: ["全球配置工具箱"],
    21: ["资金预算", "风险预算"],
    24: ["类全天候"],
    27: ["压力测试"],
    29: ["Brinson"],
    36: ["养老 FOF", "下滑曲线"],
    38: ["公募FOF", "理财FOF"],
    40: ["案例剖析"],
    42: ["Q & A"],
}


def slide_texts(slide) -> str:
    chunks = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    chunks.append(r.text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            chunks.append(r.text)
    return " | ".join(c for c in chunks if c)


def main() -> int:
    prs = Presentation(PPT)
    n = len(prs.slides)
    print(f"[verify] file = {PPT.relative_to(PROJECT_ROOT)}")
    print(f"[verify] total slides = {n}")
    assert n == EXPECTED_PAGES, f"Expected {EXPECTED_PAGES} slides, got {n}"

    titles = []
    for i, slide in enumerate(prs.slides, start=1):
        text = slide_texts(slide)
        first_chars = text[:120].replace("\n", " ")
        titles.append((i, first_chars))

    failed = 0
    for page, keywords in EXPECTED_KEYWORDS.items():
        text = titles[page - 1][1]
        for kw in keywords:
            ok = kw in text
            mark = "  OK" if ok else "FAIL"
            print(f"  [{mark}] P{page}: '{kw}' in: {text[:80]} ...")
            if not ok:
                failed += 1

    if failed:
        print(f"[verify] {failed} keyword check(s) failed.")
        return 1

    print("[verify] all checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
